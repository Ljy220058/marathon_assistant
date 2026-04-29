import argparse
import json
import pickle
import os
import time
import shutil
import logging
from pathlib import Path
from typing import List, Dict, Any, Tuple

# 配置日志
logger = logging.getLogger("vector_kb")

from marathon_qa_assistant.core.app_state import BASE_DIR, LEGACY_UPLOAD_DOCS_DIR, UPLOAD_DOCS_DIR
from marathon_qa_assistant.services.document_preprocess import normalize_text

# 引入 LangChain 和 FAISS
from langchain_community.vectorstores import FAISS
from langchain_ollama import OllamaEmbeddings
from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}

# 检索增强关键词映射
QUERY_HINTS = {
    "周期": "periodization block periodized training cycle macrocycle mesocycle",
    "营养": "nutrition marathon hydration carbohydrate fueling diet protein",
    "饮食": "diet nutrition food eating meal plan calorie",
    "补给": "fueling hydration carbohydrate electrolyte intake",
    "过度训练": "overtraining overreaching excessive loading fatigue injury",
    "伤": "injury recovery rehabilitation pain therapist orthopedic",
    "痛": "pain injury recovery rehabilitation therapist",
    "康复": "recovery rehabilitation injury physical therapy",
    "高温": "heat stress hot environment hydration safety thermal",
    "力量": "strength training resistance program gym weightlifting",
    "耐力": "endurance training aerobic running cardiovascular",
    "马拉松": "marathon running training plan race long distance",
    "配速": "pace speed velocity timing splits",
    "训练": "training exercise workout program",
    "安全": "safety risk warning caution advice health emergency",
    "风险": "risk danger safety hazard warning",
    "禁忌": "contraindication warning caution safety avoid",
    "知识库": "knowledge base scope domain expert area content coverage",
    "领域": "knowledge base scope domain expert area content coverage",
    "文档": "book document source reference literature paper pdf",
    "参考": "reference source bibliography citations documentation",
    "原则": "principle methodology theory fundamental framework guidelines",
}

DEFAULT_TEST_QUESTIONS = [
    "如何定义并安排周期化训练以支持长期跑步表现提升？",
    "马拉松训练期如何进行营养与补给规划？",
    "过度训练常见风险信号有哪些，如何识别？",
    "高温环境下训练有哪些安全建议？",
    "力量训练与耐力训练如何在同一训练计划中平衡？",
]

OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")
EMBEDDING_MODEL = "nomic-embed-text" # 可以根据实际安装的模型替换


def _build_source_metadata(file_path: Path) -> dict:
    return {
        "source_file": file_path.name,
        "source_path": str(file_path.absolute()),
    }


def candidate_source_dirs() -> list[Path]:
    return [
        UPLOAD_DOCS_DIR,
        LEGACY_UPLOAD_DOCS_DIR,
        BASE_DIR / "domain_docs",
    ]


def infer_source_path(source_file: str) -> str:
    if not source_file or source_file == "unknown":
        return ""
    for candidate_dir in candidate_source_dirs():
        candidate_path = candidate_dir / source_file
        if candidate_path.exists():
            return str(candidate_path.absolute())
    return ""


def _normalize_chunk_source(chunk: dict) -> dict:
    normalized = dict(chunk)
    source_file = str(normalized.get("source_file", "") or "").strip()
    source_path = str(normalized.get("source_path", "") or "").strip()
    if not source_path and source_file:
        normalized["source_path"] = infer_source_path(source_file)
    return normalized

def get_embeddings():
    return OllamaEmbeddings(model=EMBEDDING_MODEL, base_url=OLLAMA_BASE_URL)

def extract_pdf_pages(file_path: Path) -> list[tuple[int, str]]:
    """提取 PDF 每一页的内容"""
    try:
        import fitz
    except Exception:
        fitz = None
    if fitz is not None:
        pages = []
        with fitz.open(file_path) as doc:
            for idx, page in enumerate(doc, start=1):
                page_text = page.get_text("text") or ""
                pages.append((idx, page_text))
        return pages
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError("未安装可用的 PDF 文本抽取库，请安装 PyMuPDF 或 pypdf") from exc
    reader = PdfReader(str(file_path))
    pages = []
    for idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        pages.append((idx, page_text))
    return pages

def extract_docx_text(file_path: Path) -> str:
    """提取 docx 文件内容"""
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("处理 docx 需要安装 python-docx") from exc
    document = Document(str(file_path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

def extract_image_text_sync(file_path: Path) -> str:
    """同步方式从图片中提取文字（OCR）"""
    import asyncio
    from marathon_qa_assistant.services.ocr_service import extract_text_from_image
    try:
        loop = asyncio.get_event_loop()
        if loop.is_running():
            import concurrent.futures
            with concurrent.futures.ThreadPoolExecutor() as pool:
                future = pool.submit(lambda: asyncio.run(extract_text_from_image(str(file_path))))
                result = future.result(timeout=120)
        else:
            result = asyncio.run(extract_text_from_image(str(file_path)))
    except RuntimeError:
        result = asyncio.run(extract_text_from_image(str(file_path)))
    
    ocr_quality = result.get("ocr_quality", {})
    engine = result.get("engine", "unknown")
    text = result.get("text", "")
    
    if not text.strip():
        raise RuntimeError(f"图片未识别到有效文字 (引擎: {engine}, 评分: {ocr_quality.get('overall_score', 0)})")
    
    header = f"[OCR来源: {file_path.name} | 引擎: {engine} | 置信度: {ocr_quality.get('avg_confidence', 0):.2f}]\n"
    return header + text

def load_pages(file_path: Path) -> list[tuple[int, str]]:
    """根据扩展名加载文件页内容（支持文档和图片OCR）"""
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf_pages(file_path)
    if ext in {".txt", ".md"}:
        return [(1, file_path.read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".docx":
        return [(1, extract_docx_text(file_path))]
    if ext in IMAGE_EXTENSIONS:
        return [(1, extract_image_text_sync(file_path))]
    raise ValueError(f"不支持的文件格式: {file_path.name}")

def split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    """将文本切分为分片"""
    if chunk_size <= 0:
        raise ValueError("chunk_size 必须大于 0")
    if chunk_overlap < 0:
        raise ValueError("chunk_overlap 不能小于 0")
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap 必须小于 chunk_size")
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        start = end - chunk_overlap
    return chunks

def collect_chunks(input_dir: Path | List[Path], chunk_size: int, chunk_overlap: int) -> tuple[list[dict], list[dict]]:
    """从目录或文件列表中收集所有文本分片"""
    all_chunks = []
    file_stats = []
    
    # 兼容单个目录或文件列表
    if isinstance(input_dir, list):
        files_to_process = sorted(input_dir, key=lambda p: p.name.lower())
    else:
        if not input_dir.exists():
            return [], []
        files_to_process = sorted(input_dir.iterdir(), key=lambda p: p.name.lower())

    for file_path in files_to_process:
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        page_count = 0
        chunk_count = 0
        source_meta = _build_source_metadata(file_path)
        try:
            pages = load_pages(file_path)
            page_count = len(pages)
            for page_num, raw_page in pages:
                cleaned_page = normalize_text(raw_page)
                page_chunks = split_text(cleaned_page, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
                for idx, chunk_text in enumerate(page_chunks, start=1):
                    chunk_id = f"{file_path.stem}_p{page_num:04d}_c{idx:04d}"
                    all_chunks.append(
                        {
                            "chunk_id": chunk_id,
                            "source_file": source_meta["source_file"],
                            "source_path": source_meta["source_path"],
                            "page": page_num,
                            "text": chunk_text,
                        }
                    )
                chunk_count += len(page_chunks)
            file_stats.append(
                {
                    "source_file": source_meta["source_file"],
                    "source_path": source_meta["source_path"],
                    "pages": page_count,
                    "chunks": chunk_count,
                }
            )
        except Exception as exc:
            file_stats.append(
                {
                    "source_file": source_meta["source_file"],
                    "source_path": source_meta["source_path"],
                    "pages": page_count,
                    "chunks": chunk_count,
                    "error": str(exc),
                }
            )
    return all_chunks, file_stats

def build_hybrid_indices(chunks: list[dict]):
    """
    为了兼容 module_kb.py 的签名，返回三个占位符。
    真正的向量构建在 save_outputs 中执行。
    """
    return "faiss_vectorizer", "faiss_matrix", "faiss_bm25"

def save_outputs(output_dir: Path, chunks: list[dict], vectorizer, matrix, bm25) -> dict:
    """保存构建好的向量库和分片数据"""
    # 确保 output_dir 及其所有父目录都存在
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        logger.error(f"创建输出目录失败: {e}")
        # 尝试强制创建父目录
        os.makedirs(str(output_dir), exist_ok=True)
    
    # 1. 构建 FAISS 向量库
    faiss_dir = output_dir / "faiss_db"
    
    if faiss_dir.exists():
        try:
            shutil.rmtree(faiss_dir)
            time.sleep(0.5)
        except Exception as e:
            logger.warning(f"清理旧目录失败: {e}")

    # 再次确保 faiss_dir 存在
    try:
        faiss_dir.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        logger.error(f"创建 FAISS 目录失败: {e}")
        os.makedirs(str(faiss_dir), exist_ok=True)
            
    embeddings = get_embeddings()
    logger.info(f"正在构建 FAISS 向量库: {faiss_dir}")
    
    # 准备 Document 对象
    docs = []
    for c in chunks:
        docs.append(Document(
            page_content=c["text"],
            metadata={
                "chunk_id": c["chunk_id"],
                "source_file": c["source_file"],
                "source_path": c.get("source_path", ""),
                "page": c["page"]
            }
        ))
        
    faiss_store = FAISS.from_documents(docs, embeddings)
    faiss_dir_str = str(faiss_dir)
    
    # 在保存时同样尝试处理 Windows 中文路径问题
    import os
    if os.name == 'nt':
        try:
            rel_path = os.path.relpath(faiss_dir_str, os.getcwd())
            faiss_store.save_local(rel_path)
            logger.info(f"FAISS 索引已保存 (相对路径: {rel_path})")
        except Exception as rel_e:
            logger.debug(f"相对路径保存失败: {rel_e}，尝试使用绝对路径")
            faiss_store.save_local(faiss_dir_str)
    else:
        faiss_store.save_local(faiss_dir_str)
    
    # 2. 保存 chunks.jsonl 供其他模块使用
    chunks_path = output_dir / "chunks.jsonl"
    with chunks_path.open("w", encoding="utf-8") as f:
        for row in chunks:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
            
    logger.info("FAISS indexing complete!")
    
    return {
        "chunks_file": str(chunks_path),
        "faiss_dir": str(faiss_dir),
    }

def load_chunks(chunks_file: Path) -> list[dict]:
    """加载 jsonl 格式的分片数据"""
    chunks = []
    with chunks_file.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            chunks.append(_normalize_chunk_source(json.loads(line)))
    return chunks

def load_vector_kb(vector_dir: Path):
    """
    加载向量库，返回 (chunks, vectorizer, matrix, bm25) 结构。
    实际上 matrix 位置放置的是 FAISS 实例。
    """
    chunks_file = vector_dir / "chunks.jsonl"
    faiss_dir = vector_dir / "faiss_db"
    
    if not chunks_file.exists():
        raise FileNotFoundError(f"未找到 chunks 文件: {chunks_file}")
    
    chunks = load_chunks(chunks_file)
    
    # 初始化 FAISS 客户端作为 "matrix"
    embeddings = get_embeddings()
    faiss_store = None
    
    # 更加严谨的可用性判断：不仅检查目录，还要检查索引文件 index.faiss
    faiss_index_file = faiss_dir / "index.faiss"
    
    if faiss_index_file.exists():
        try:
            import os
            # 在 Windows 下，避免 resolve() 穿透 Junction/Symlink，直接使用绝对路径字符串
            faiss_dir_str = str(faiss_dir.absolute())
            
            # 解决 FAISS C++ 底层 fopen 不支持 Windows 中文绝对路径的问题
            # 方案1: 获取 Windows 8.3 短路径 (仅限 ASCII)
            short_path = faiss_dir_str
            if os.name == 'nt':
                try:
                    import ctypes
                    from ctypes import wintypes
                    _GetShortPathNameW = ctypes.windll.kernel32.GetShortPathNameW
                    _GetShortPathNameW.argtypes = [wintypes.LPCWSTR, wintypes.LPWSTR, wintypes.DWORD]
                    _GetShortPathNameW.restype = wintypes.DWORD
                    
                    buf_size = _GetShortPathNameW(faiss_dir_str, None, 0)
                    if buf_size > 0:
                        buf = ctypes.create_unicode_buffer(buf_size)
                        if _GetShortPathNameW(faiss_dir_str, buf, buf_size):
                            short_path = buf.value
                except Exception as e:
                    logger.debug(f"获取短路径失败: {e}")
            
            # 方案2: 相对路径 (通常不含项目根目录的中文)
            try:
                rel_path = os.path.relpath(faiss_dir_str, os.getcwd())
            except ValueError:
                rel_path = faiss_dir_str

            # 依次尝试 短路径 -> 相对路径 -> 绝对路径
            loaded = False
            for try_path in [short_path, rel_path, faiss_dir_str]:
                try:
                    faiss_store = FAISS.load_local(try_path, embeddings, allow_dangerous_deserialization=True)
                    logger.info(f"FAISS 索引加载成功 (使用路径: {try_path})")
                    loaded = True
                    break
                except Exception as e:
                    logger.debug(f"尝试路径 {try_path} 失败: {e}")
            
            if not loaded:
                # 方案3: 内存反序列化兜底 (完全绕过 FAISS C++ 文件读取)
                logger.warning("所有常规路径加载 FAISS 失败，尝试内存反序列化兜底加载...")
                import faiss
                import pickle
                import numpy as np
                
                with open(faiss_index_file, "rb") as f:
                    index_bytes = f.read()
                
                with open(faiss_dir / "index.pkl", "rb") as f:
                    docstore, index_to_docstore_id = pickle.load(f)
                    
                try:
                    index_array = np.frombuffer(index_bytes, dtype=np.uint8)
                    index = faiss.deserialize_index(index_array)
                except Exception:
                    index = faiss.deserialize_index(index_bytes)
                    
                faiss_store = FAISS(
                    embedding_function=embeddings.embed_query,
                    index=index,
                    docstore=docstore,
                    index_to_docstore_id=index_to_docstore_id
                )
                logger.info("FAISS 索引加载成功(内存反序列化)")

        except Exception as e:
            logger.error(f"加载 FAISS 失败 (路径: {faiss_dir}): {e}")
            faiss_store = None
    else:
        logger.warning(f"FAISS 索引文件缺失: {faiss_index_file}，将以空库运行。")
        
    # 返回: chunks, vectorizer(dummy), matrix(faiss), bm25(dummy)
    return chunks, "faiss_vectorizer", faiss_store, "faiss_bm25"

def retrieve(question: str, chunks: list[dict], vectorizer, matrix, top_k: int, bm25=None) -> list[dict]:
    """
    执行检索逻辑。这里的 matrix 实际上是 FAISS 实例。
    """
    faiss_store = matrix
    if not faiss_store:
        logger.warning("FAISS 数据库未初始化，返回空检索结果。")
        return []
        
    query = normalize_text(question)
    enhanced_query = query
    # 注入检索提示词增强效果
    for k, v in QUERY_HINTS.items():
        if k in query:
            enhanced_query = f"{enhanced_query} {v}"
            
    # 执行 FAISS 向量相似度检索 (带分数)
    try:
        results = faiss_store.similarity_search_with_score(enhanced_query, k=top_k)
    except Exception as e:
        logger.error(f"FAISS 检索失败: {e}")
        return []
        
    hits = []
    for doc, distance in results:
        # FAISS 默认使用 L2 距离，转换为相似度得分
        # 得分转换公式：score = 1.0 / (1.0 + distance)
        score = round(1.0 / (1.0 + float(distance)), 6)
        
        hits.append({
            "score": score,
            "chunk_id": doc.metadata.get("chunk_id", "unknown"),
            "source_file": doc.metadata.get("source_file", "unknown"),
            "source_path": doc.metadata.get("source_path", "") or infer_source_path(doc.metadata.get("source_file", "")),
            "page": doc.metadata.get("page", 1),
            "text": doc.page_content,
            "distance": float(distance)
        })
        
    # 按 score 降序排列
    hits = sorted(hits, key=lambda x: x["score"], reverse=True)
    return hits

def build_answer(hits: list[dict]) -> str:
    """从检索结果构建简短回答片段（供调试/测试用）"""
    if not hits:
        return "未检索到有效证据，请尝试更具体的问题。"
    snippets = []
    for hit in hits[:3]:
        text = hit["text"].replace("\n", " ").strip()
        snippets.append(text[:220])
    return " ".join(snippets).strip()

def load_test_questions(questions_file: Path | None) -> list[str]:
    """加载测试问题列表"""
    if questions_file is None:
        return DEFAULT_TEST_QUESTIONS
    if not questions_file.exists():
        raise FileNotFoundError(f"测试问题文件不存在: {questions_file}")
    if questions_file.suffix.lower() == ".json":
        data = json.loads(questions_file.read_text(encoding="utf-8"))
        if not isinstance(data, list):
            raise ValueError("JSON 问题文件应为字符串数组")
        questions = [str(x).strip() for x in data if str(x).strip()]
        if not questions:
            raise ValueError("JSON 问题文件为空")
        return questions
    questions = []
    for line in questions_file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line:
            questions.append(line)
    if not questions:
        raise ValueError("问题文件为空")
    return questions

def run_rag_test(vector_dir: Path, output_file: Path, top_k: int, questions: list[str]) -> dict:
    """运行 RAG 检索基准测试"""
    chunks, vectorizer, matrix, bm25 = load_vector_kb(vector_dir)
    results = []
    for question in questions:
        hits = retrieve(question, chunks, vectorizer, matrix, top_k=top_k, bm25=bm25)
        citations = [
            {
                "source_file": x["source_file"],
                "page": x["page"],
                "chunk_id": x["chunk_id"],
                "score": x["score"],
            }
            for x in hits
        ]
        results.append(
            {
                "question": question,
                "answer": build_answer(hits),
                "citations": citations,
                "retrieved_chunks": hits,
            }
        )
    report = {
        "vector_dir": str(vector_dir),
        "top_k": top_k,
        "question_count": len(questions),
        "results": results,
    }
    output_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["build", "test"], default="build")
    parser.add_argument("--input-dir", default="domain_docs")
    parser.add_argument("--output-dir", default="vector_kb")
    parser.add_argument("--report-file", default="vector_kb_report.json")
    parser.add_argument("--chunk-size", type=int, default=500)
    parser.add_argument("--chunk-overlap", type=int, default=50)
    parser.add_argument("--vector-dir", default="vector_kb")
    parser.add_argument("--test-output", default="rag_test_results.json")
    parser.add_argument("--top-k", type=int, default=5)
    parser.add_argument("--questions-file", default=None)
    args = parser.parse_args()

    if args.mode == "build":
        logger.info("Starting build mode...")
        input_dir = Path(args.input_dir).absolute()
        output_dir = Path(args.output_dir).absolute()
        report_file = Path(args.report_file).absolute()
        if not input_dir.exists():
            raise FileNotFoundError(f"输入目录不存在: {input_dir}")
        
        chunks, file_stats = collect_chunks(
            input_dir=input_dir,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
        )
        if not chunks:
            raise RuntimeError("未生成任何 chunk，无法构建向量库")
            
        vectorizer, matrix, bm25 = build_hybrid_indices(chunks)
        saved_paths = save_outputs(output_dir, chunks, vectorizer, matrix, bm25)
        
        report = {
            "input_dir": str(input_dir),
            "output_dir": str(output_dir),
            "chunk_size": args.chunk_size,
            "chunk_overlap": args.chunk_overlap,
            "total_chunks": len(chunks),
            "engine": "FAISS + Nomic Embedding",
            "files": file_stats,
            "artifacts": saved_paths,
        }
        report_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        logger.info("Build finished successfully.")
        return

    vector_dir = Path(args.vector_dir).absolute()
    output_file = Path(args.test_output).absolute()
    questions_file = Path(args.questions_file).absolute() if args.questions_file else None
    questions = load_test_questions(questions_file)
    test_report = run_rag_test(vector_dir=vector_dir, output_file=output_file, top_k=args.top_k, questions=questions)
    print(json.dumps(test_report, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
